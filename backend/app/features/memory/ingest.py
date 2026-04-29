"""File ingestion — parse uploaded documents and split them into chunks.

Supported types: .txt, .md, .pdf, .docx
Hard limits:
- max file size: 10 MB
- max chunks per file: 500

PDF and DOCX parsing are sync (pypdf / python-docx) and CPU/IO-bound, so they
run inside `asyncio.to_thread`.
"""

import asyncio
import io
import logging
from pathlib import PurePosixPath

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

logger = logging.getLogger(__name__)

MAX_FILE_SIZE_BYTES: int = 10 * 1024 * 1024  # 10 MB
MAX_CHUNKS_PER_FILE: int = 500
ALLOWED_EXTENSIONS: frozenset[str] = frozenset({".txt", ".md", ".pdf", ".docx"})

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=800,
    chunk_overlap=100,
    separators=["\n\n", "\n", ". ", " ", ""],
)


class UnsupportedFileTypeError(ValueError):
    pass


class FileTooLargeError(ValueError):
    pass


def _parse_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _parse_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    # Also extract text from tables — commonly contains key info in assignment docs.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def _parse_text(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def _extract_text_sync(filename: str, content: bytes) -> str:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{suffix}'. Allowed: {sorted(ALLOWED_EXTENSIONS)}"
        )
    if suffix == ".pdf":
        return _parse_pdf(content)
    if suffix == ".docx":
        return _parse_docx(content)
    return _parse_text(content)


async def extract_text(filename: str, content: bytes) -> str:
    """Async wrapper around the sync parser; offloads PDF parsing off the loop."""
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise FileTooLargeError(
            f"File exceeds {MAX_FILE_SIZE_BYTES} bytes (got {len(content)})"
        )
    return await asyncio.to_thread(_extract_text_sync, filename, content)


def chunk_text(text: str) -> list[str]:
    """Split text into ~800-char overlapping chunks. Caps at MAX_CHUNKS_PER_FILE."""
    chunks = _splitter.split_text(text)
    if len(chunks) > MAX_CHUNKS_PER_FILE:
        logger.warning(
            "File produced %s chunks; truncating to %s", len(chunks), MAX_CHUNKS_PER_FILE
        )
        chunks = chunks[:MAX_CHUNKS_PER_FILE]
    return chunks
